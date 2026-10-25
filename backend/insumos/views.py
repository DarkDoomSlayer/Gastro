from django.db.models import Q
import openpyxl
import zipfile
import re
from decimal import Decimal
from rest_framework import viewsets, filters, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from .models import Insumo
from .serializers import InsumoSerializer

from rest_framework.authentication import SessionAuthentication, BasicAuthentication

class CsrfExemptSessionAuthentication(SessionAuthentication):
    def enforce_csrf(self, request):
        return  # Deshabilita la verificación CSRF estricta para APIs tipo SPA

class SanitizedSearchFilter(filters.SearchFilter):
    """Filtro de búsqueda que sanitiza y limita caracteres para prevenir ataques de saturación o sintaxis maliciosa."""
    def get_search_terms(self, request):
        params = super().get_search_terms(request)
        sanitized = []
        for term in params:
            # Eliminar caracteres de control peligrosos, comillas escapadas y limitar a 80 caracteres
            clean_term = re.sub(r'[\x00-\x1f\x7f<>;\"\'%\\]', '', term).strip()[:80]
            if clean_term:
                sanitized.append(clean_term)
        return sanitized

class InsumoViewSet(viewsets.ModelViewSet):
    authentication_classes = [CsrfExemptSessionAuthentication, BasicAuthentication]
    serializer_class = InsumoSerializer
    pagination_class = None
    parser_classes = [JSONParser, MultiPartParser, FormParser]
    filter_backends = [SanitizedSearchFilter, filters.OrderingFilter]
    search_fields = ['nombre', 'tipo', 'unidad']
    ordering_fields = ['id', 'nombre', 'tipo', 'precio']
    ordering = ['id']

    def get_queryset(self):
        user = self.request.user
        if user.is_authenticated:
            # Devuelve únicamente los insumos propios del usuario autenticado
            return Insumo.objects.filter(usuario=user, activo=True)
        return Insumo.objects.none()

    def create(self, request, *args, **kwargs):
        nombre = request.data.get('nombre', '').strip()
        if not nombre:
            return Response({'error': 'El nombre del ingrediente es requerido'}, status=status.HTTP_400_BAD_REQUEST)
        user = request.user if request.user.is_authenticated else None
        
        # Validar si ya existe un insumo con ese nombre (ignora mayúsculas/minúsculas)
        filtro = {'nombre__iexact': nombre, 'activo': True}
        if user:
            filtro['usuario'] = user
        else:
            filtro['usuario__isnull'] = True

        if Insumo.objects.filter(**filtro).exists():
            return Response({'error': f'Ya existe un ingrediente registrado con el nombre "{nombre}".'}, status=status.HTTP_400_BAD_REQUEST)

        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        if user:
            serializer.save(usuario=user)
        else:
            serializer.save()
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        forzar = request.query_params.get('forzar', 'false').lower() == 'true'
        usos = instance.usos_en_recetas.all()
        recetas_nombres = list(set(uso.receta.nombre for uso in usos if uso.receta))

        if usos.exists() and not forzar:
            return Response({
                'bloqueado': True,
                'mensaje': f'El insumo "{instance.nombre}" no se puede eliminar porque se encuentra registrado en {len(recetas_nombres)} receta(s): {", ".join(recetas_nombres)}.',
                'recetas': recetas_nombres
            }, status=status.HTTP_409_CONFLICT)

        # Si el usuario confirma forzar, primero se retira de esas recetas
        if usos.exists() and forzar:
            usos.delete()

        self.perform_destroy(instance)
        return Response(status=status.HTTP_204_NO_CONTENT)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', False)
        instance = self.get_object()
        forzar = request.data.get('forzar', False) or request.query_params.get('forzar', 'false').lower() == 'true'
        usos = instance.usos_en_recetas.all()
        recetas_nombres = list(set(uso.receta.nombre for uso in usos if uso.receta))

        nuevo_nombre = request.data.get('nombre', '').strip()
        if nuevo_nombre and nuevo_nombre.lower() != instance.nombre.lower():
            user = request.user if request.user.is_authenticated else None
            filtro = {'nombre__iexact': nuevo_nombre, 'activo': True}
            if user:
                filtro['usuario'] = user
            else:
                filtro['usuario__isnull'] = True

            if Insumo.objects.filter(**filtro).exclude(id=instance.id).exists():
                return Response({'error': f'Ya existe otro ingrediente registrado con el nombre "{nuevo_nombre}".'}, status=status.HTTP_400_BAD_REQUEST)

        if usos.exists() and not forzar:
            return Response({
                'bloqueado': True,
                'mensaje': f'El insumo "{instance.nombre}" no se puede editar directamente porque se encuentra en {len(recetas_nombres)} receta(s): {", ".join(recetas_nombres)}.',
                'recetas': recetas_nombres
            }, status=status.HTTP_409_CONFLICT)

        # Si el usuario confirma forzar para editar, primero se retira de esas recetas
        if usos.exists() and forzar:
            usos.delete()

        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response(serializer.data)

    @action(detail=False, methods=['post'], url_path='importar-excel')
    def importar_excel(self, request):
        """
        Importa ingredientes desde la hoja 'BASE DE DATOS' de un archivo Excel.
        Evita registros duplicados por nombre.
        Columnas esperadas:
        - Col A: Ingrediente
        - Col B: Tipo de ingrediente (Materia prima / Complementaria)
        - Col C: Presentación
        - Col D: Unidad (Kg / Lt)
        - Col E: Precio
        """
        archivo = request.FILES.get('file')
        if not archivo:
            return Response({'error': 'No se proporcionó ningún archivo Excel (.xlsx)'}, status=status.HTTP_400_BAD_REQUEST)

        # 1. Validación de tamaño máximo (5MB) para prevenir ataques de agotamiento de memoria
        MAX_EXCEL_BYTES = 5 * 1024 * 1024
        if archivo.size > MAX_EXCEL_BYTES:
            return Response({'error': 'El archivo excede el límite permitido de 5 MB.'}, status=status.HTTP_400_BAD_REQUEST)

        # 2. Validación de Magic Bytes de archivo ZIP/OpenXML (50 4B 03 04 -> PK\x03\x04)
        header = archivo.read(4)
        archivo.seek(0)
        if header != b'PK\x03\x04':
            return Response({
                'error': 'El archivo no es un documento de cálculo Excel válido (firma binaria / magic bytes inválidos).'
            }, status=status.HTTP_400_BAD_REQUEST)

        # 3. Validación de estructura interna de archivo OpenXML (.xlsx)
        try:
            if not zipfile.is_zipfile(archivo):
                return Response({'error': 'El archivo no tiene una estructura de compresión válida.'}, status=status.HTTP_400_BAD_REQUEST)
            archivo.seek(0)
            with zipfile.ZipFile(archivo, 'r') as zf:
                namelist = zf.namelist()
                # Un Excel .xlsx debe contener al menos el archivo de workbook o tipos de contenido
                tiene_workbook = any('xl/workbook' in name or '[Content_Types].xml' in name for name in namelist)
                if not tiene_workbook:
                    return Response({
                        'error': 'El archivo comprimido no contiene la estructura esperada de un libro de Excel.'
                    }, status=status.HTTP_400_BAD_REQUEST)
        except Exception:
            return Response({'error': 'Error al inspeccionar la estructura binaria del archivo.'}, status=status.HTTP_400_BAD_REQUEST)

        archivo.seek(0)
        try:
            wb = openpyxl.load_workbook(archivo, data_only=True)
            # Buscar hoja BASE DE DATOS de manera flexible
            sheet_name = None
            for name in wb.sheetnames:
                if 'BASE' in name.upper() and 'DATO' in name.upper():
                    sheet_name = name
                    break
            
            if not sheet_name:
                sheet_name = wb.sheetnames[0]  # Tomar la primera hoja si no coincide exactamente

            ws = wb[sheet_name]

            creados = 0
            actualizados = 0
            omitidos = 0

            # Iterar filas a partir de la fila 2 (la fila 1 suelen ser encabezados)
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue

                nombre_raw = str(row[0]).strip()
                if not nombre_raw or nombre_raw.startswith('#') or nombre_raw.upper() in ['INGREDIENTE', 'BASE DE DATOS', 'TOTAL']:
                    continue

                # Tipo de ingrediente
                tipo_raw = str(row[1]).strip().lower() if len(row) > 1 and row[1] is not None else 'materia prima'
                tipo_val = 'complementaria' if 'comp' in tipo_raw else 'materia_prima'

                # Presentación (si tiene #REF!, #ERROR! o está vacío, dejar en None)
                presentacion_val = None
                if len(row) > 2 and row[2] is not None:
                    val_str = str(row[2]).strip()
                    if not val_str.startswith('#') and val_str.upper() not in ['NONE', 'NULL', '']:
                        try:
                            presentacion_val = Decimal(val_str.replace(',', ''))
                        except Exception:
                            presentacion_val = None

                # Unidad (Kg o Lt)
                unidad_raw = str(row[3]).strip() if len(row) > 3 and row[3] is not None else 'Kg'
                unidad_val = 'Lt' if 'lt' in unidad_raw.lower() or 'l' == unidad_raw.lower() else 'Kg'

                # Precio (si tiene #REF!, #ERROR! o está vacío, dejar en None)
                precio_val = None
                if len(row) > 4 and row[4] is not None:
                    val_str = str(row[4]).replace('$', '').replace(',', '').strip()
                    if not val_str.startswith('#') and val_str.upper() not in ['NONE', 'NULL', '']:
                        try:
                            precio_val = Decimal(val_str)
                        except Exception:
                            precio_val = None

                user = request.user if request.user.is_authenticated else None

                # No duplicar por nombre dentro del catálogo del propio usuario
                filtro_busqueda = {'nombre__iexact': nombre_raw}
                if user:
                    filtro_busqueda['usuario'] = user
                else:
                    filtro_busqueda['usuario__isnull'] = True

                insumo = Insumo.objects.filter(**filtro_busqueda).first()
                if not insumo:
                    Insumo.objects.create(
                        usuario=user,
                        nombre=nombre_raw,
                        tipo=tipo_val,
                        presentacion=presentacion_val,
                        unidad=unidad_val,
                        precio=precio_val,
                        activo=True,
                    )
                    creados += 1
                else:
                    # Si ya existe para este usuario, actualizamos sus datos sin duplicar
                    insumo.tipo = tipo_val
                    insumo.presentacion = presentacion_val
                    insumo.unidad = unidad_val
                    insumo.precio = precio_val
                    insumo.activo = True
                    insumo.save()
                    actualizados += 1

            return Response({
                'mensaje': 'Importación procesada exitosamente',
                'hoja_procesada': sheet_name,
                'creados': creados,
                'actualizados': actualizados,
                'total': creados + actualizados,
            })

        except Exception as e:
            return Response({'error': f'Error al procesar el archivo Excel: {str(e)}'}, status=status.HTTP_400_BAD_REQUEST)
